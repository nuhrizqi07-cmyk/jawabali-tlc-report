"""
Generator DOCX laporan harian TLC dari entry parser.
Output: bytes docx siap download.
"""

import io
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CABANG_NAMA = {
    'PWO': 'PASURUAN (PWO)',
    'CXP': 'CXP',
    'KBM': 'KBM',
    'BJN': 'BANJARNEGARA (BJN)',
}

# Warna institusional
BLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x33, 0x33, 0x33)
GRAY = RGBColor(0x66, 0x66, 0x66)


def _fmt_tanggal(dmy):
    """'11/08/26' -> '11 Agustus 2026'"""
    bulan = ['Januari', 'Februari', 'Maret', 'April', 'Mei', 'Juni',
             'Juli', 'Agustus', 'September', 'Oktober', 'November', 'Desember']
    try:
        d, mo, y = dmy.split('/')
        return f"{int(d)} {bulan[int(mo)-1]} 20{y}"
    except Exception:
        return dmy


def _set_cell_bg(cell, hex_color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def build_docx(entries, tanggal, cabangs=None, judul=None):
    """Buat docx laporan harian. entries sudah difilter per tanggal (opsional cabang)."""
    doc = Document()

    # margin agak sempit biar muat
    for section in doc.sections:
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)

    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10)

    # ===== Header =====
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('LAPORAN HARIAN TLC')
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Grup NETWORK JAWABALI')
    r.font.size = Pt(12)
    r.font.color.rgb = DARK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Periode: {_fmt_tanggal(tanggal)}")
    r.font.size = Pt(11)
    r.font.color.rgb = GRAY

    if judul:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(judul)
        r.italic = True
        r.font.size = Pt(10)
        r.font.color.rgb = GRAY

    doc.add_paragraph()

    # ===== Ringkasan =====
    if cabangs is None:
        cabangs = ['PWO', 'CXP', 'KBM', 'BJN']
    cabangs = [c.upper() for c in cabangs]

    h = doc.add_paragraph()
    r = h.add_run('RINGKASAN')
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = BLUE

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    for i, t in enumerate(['Cabang', 'TLC', 'Eskalasi']):
        hdr[i].text = t
        for para in hdr[i].paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = True
    _set_cell_bg(hdr[0], 'D9E2F3')
    _set_cell_bg(hdr[1], 'D9E2F3')
    _set_cell_bg(hdr[2], 'D9E2F3')

    for c in cabangs:
        items = [e for e in entries if e['cabang'] == c]
        n_tlc = sum(1 for e in items if e['tipe'] == 'tlc')
        n_esc = sum(1 for e in items if e['tipe'] == 'eskalasi')
        row = table.add_row().cells
        row[0].text = CABANG_NAMA.get(c, c)
        row[1].text = str(n_tlc)
        row[2].text = str(n_esc)
        for cell in row:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # ===== Detail per cabang =====
    for c in cabangs:
        items = [e for e in entries if e['cabang'] == c]
        if not items:
            continue
        h = doc.add_paragraph()
        r = h.add_run(f"{CABANG_NAMA.get(c, c)}  ({len(items)} item)")
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = BLUE

        for idx, e in enumerate(items, 1):
            # Item header
            p = doc.add_paragraph()
            r = p.add_run(f"{idx}. AWB: {e['awb'] or '-'}  [{e['tipe'].upper()}]")
            r.bold = True
            r.font.size = Pt(10)

            def _field(label, val):
                if not val:
                    return
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25)
                r = p.add_run(f"{label}: ")
                r.bold = True
                r.font.size = Pt(10)
                r2 = p.add_run(val)
                r2.font.size = Pt(10)

            if e['tipe'] == 'tlc':
                _field('Chat Pengirim', e['chat'])
                _field('Ket', e['ket'])
            else:
                _field('Client', e['client'])
                _field('Kendala', e['kendala'])
                _field('Butuh Bantuan', e['butuh_bantuan'])
                _field('Patokan', e['patokan'])
            _field('Waktu', e['waktu'])

    # ===== Footer =====
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('— Laporan dibuat otomatis dari backup chat WhatsApp —')
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def filename_for(tanggal):
    d, mo, y = tanggal.split('/')
    return f"Laporan_Harian_TLC_{d}{mo}{y}.docx"
